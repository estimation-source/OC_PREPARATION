from __future__ import annotations

import io
import math
import os
import re
import sys
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
from PIL import Image
import pypdf
import streamlit as st

# ============================================================
# 1. Streamlit Page Config
# ============================================================
st.set_page_config(
    page_title="WIN-SQUARE | OC PREPARATION Engine",
    layout="wide",
    page_icon="🪟",
    initial_sidebar_state="expanded",
)

# ============================================================
# 2. Header and Sidebar Toggle Button Styling
# ============================================================
st.markdown(
    """
    <style>
    header[data-testid="stHeader"] {
        z-index: 99999 !important;
        background: transparent !important;
    }

    button[data-testid="stSidebarCollapsedControl"],
    button[data-testid="stSidebarNavCollapseButton"] {
        display: flex !important;
        visibility: visible !important;
        opacity: 1 !important;
        background-color: #FF4B4B !important;
        color: white !important;
        border-radius: 8px !important;
        position: fixed !important;
        top: 12px !important;
        left: 12px !important;
        z-index: 999999 !important;
        box-shadow: 0px 3px 8px rgba(0,0,0,0.3) !important;
    }

    button[data-testid="stSidebarCollapsedControl"] svg,
    button[data-testid="stSidebarNavCollapseButton"] svg {
        fill: white !important;
        color: white !important;
        width: 22px !important;
        height: 22px !important;
    }

    [data-testid="stStatusWidget"],
    #MainMenu, 
    footer {
        display: none !important;
        visibility: hidden !important;
    }
    </style>
""",
    unsafe_allow_html=True,
)

# ============================================================
# 3. UI Layout & Fonts CSS
# ============================================================
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Inter', sans-serif;
        background-color: #f4f6f9;
        color: #334155;
    }

    .main .block-container {
        padding-top: 1.5rem;
        padding-bottom: 2rem;
        max-width: 98%;
    }

    [data-testid="stSidebar"] {
        background-color: #f1f5f9;
        border-right: 1px solid #e2e8f0;
    }
    
    .quick-guide-title {
        font-size: 15px;
        font-weight: 700;
        color: #0f172a;
        margin-top: 15px;
        margin-bottom: 12px;
    }
    
    .quick-guide-step {
        font-size: 13px;
        color: #475569;
        margin-bottom: 10px;
        line-height: 1.4;
    }

    .hero-container {
        background: #ffffff;
        border-radius: 16px;
        padding: 24px 30px;
        border: 1px solid #e2e8f0;
        box-shadow: 0 2px 4px rgba(0,0,0,0.02);
        margin-bottom: 24px;
        display: flex;
        justify-content: space-between;
        align-items: center;
    }

    .hero-title-text {
        font-size: 22px;
        font-weight: 800;
        color: #0f172a;
        margin: 0;
    }

    .hero-sub-text {
        font-size: 13px;
        color: #64748b;
        margin-top: 4px;
    }

    .step-title {
        font-size: 16px;
        font-weight: 700;
        color: #1e293b;
        margin-bottom: 12px;
        display: flex;
        align-items: center;
        gap: 8px;
    }

    div.stButton > button[kind="primary"] {
        background-color: #2563eb !important;
        background: #2563eb !important;
        border: 1px solid #2563eb !important;
        color: #ffffff !important;
        border-radius: 6px !important;
        height: 38px !important;
        padding: 0 16px !important;
        box-shadow: 0 1px 2px rgba(37, 99, 235, 0.2) !important;
    }

    div.stDownloadButton > button {
        background-color: #059669 !important;
        background: #059669 !important;
        border: 1px solid #059669 !important;
        color: #ffffff !important;
        border-radius: 6px !important;
        height: 38px !important;
        padding: 0 16px !important;
        box-shadow: 0 1px 2px rgba(5, 150, 105, 0.2) !important;
    }

    div.stButton > button p, div.stButton > button span,
    div.stDownloadButton > button p, div.stDownloadButton > button span {
        color: #ffffff !important;
        font-weight: 500 !important;
        font-size: 13px !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


def get_image_path(filename: str) -> str:
    if hasattr(sys, "_MEIPASS"):
        return os.path.join(sys._MEIPASS, filename)
    return os.path.join(os.path.abspath("."), filename)


# =========================================================
# SIDEBAR
# =========================================================
with st.sidebar:
    logo_file = get_image_path("logo.png")
    if os.path.exists(logo_file):
        col_s1, col_s2, col_s3 = st.columns([1, 2, 1])
        with col_s2:
            st.image(Image.open(logo_file), width=110)
    else:
        st.markdown(
            "<h2 style='text-align: center; color:#1e293b;'><b>win square</b></h2>",
            unsafe_allow_html=True,
        )

    st.markdown("---")
    st.markdown(
        "<div class='quick-guide-title'>💡 Quick Guide</div>",
        unsafe_allow_html=True,
    )
    st.markdown(
        """
        <div class='quick-guide-step'><b>1.</b> Upload Quotation file (.xlsx or .pdf).</div>
        <div class='quick-guide-step'><b>2.</b> Engine extracts drawings, auto-resizes to fit sheet cleanly (Image 2 style).</div>
        <div class='quick-guide-step'><b>3.</b> Splits Quantity (e.g. W1 -> W1A, W1B...) and formats Glass Thickness.</div>
        <div class='quick-guide-step'><b>4.</b> Download final <b>MEASUREMENT</b> sheet.</div>
        """,
        unsafe_allow_html=True,
    )

# =========================================================
# HEADER HERO BANNER
# =========================================================
st.markdown(
    """
    <div class="hero-container">
        <div>
            <div class="hero-title-text">OC PREPARATION Engine</div>
            <div class="hero-sub-text">Automatic Measurement Extraction, Drawing Resizing & Alignment</div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

# ============================================================
# Helper Functions
# ============================================================

def get_code_suffix(index: int) -> str:
    """Generates A, B, C... Z, AA, AB... for quantity splitting."""
    result = ""
    while True:
        result = chr(65 + (index % 26)) + result
        index = (index // 26) - 1
        if index < 0:
            break
    return result


def extract_glass_thickness(glass_text: Any) -> str:
    """Extracts glass thickness and formats as '20 MM', '5 MM', '11.52 MM', etc."""
    if not glass_text or pd.isna(glass_text):
        return ""
    text = str(glass_text).strip()

    match = re.search(r"(\d+(?:\.\d+)?)\s*(MM|mm)", text)
    if match:
        return f"{match.group(1)} MM"

    match_lam = re.search(r"\)(\d+(?:\.\d+)?)\s*(MM|mm)?", text, re.IGNORECASE)
    if match_lam:
        return f"{match_lam.group(1)} MM"

    return ""


# ============================================================
# Parsing Logic (PDF & Excel)
# ============================================================

def parse_pdf_quotation(file_bytes: io.BytesIO) -> List[Dict[str, Any]]:
    reader = pypdf.PdfReader(file_bytes)
    extracted_items = []

    for page in reader.pages:
        page_text = page.extract_text() or ""

        page_images = []
        try:
            for img_obj in page.images:
                page_images.append(img_obj.data)
        except Exception:
            pass

        code_blocks = re.split(r"(Code\s*:\s*)", page_text, flags=re.IGNORECASE)

        for i in range(1, len(code_blocks), 2):
            block_text = code_blocks[i] + code_blocks[i + 1]

            code_match = re.search(r"Code\s*:\s*([^\n\r]+)", block_text)
            name_match = re.search(r"Name\s*:\s*([^\n\r]+)", block_text)
            loc_match = re.search(r"Location\s*:\s*([^\n\r]+)", block_text)
            size_match = re.search(
                r"Size\s*:\s*W\s*=\s*([\d\.]+);\s*H\s*=\s*([\d\.]+)", block_text
            )
            prof_match = re.search(r"Profile System\s*:\s*([^\n\r]+)", block_text)
            glass_match = re.search(r"Glass\s*:\s*([^\n\r]+)", block_text)
            qty_match = re.search(
                r"Quantity\s*:\s*(\d+)", block_text, flags=re.IGNORECASE
            )

            base_code = code_match.group(1).strip() if code_match else ""
            name = name_match.group(1).strip() if name_match else ""
            location = loc_match.group(1).strip() if loc_match else ""
            w_val = float(size_match.group(1)) if size_match else ""
            h_val = float(size_match.group(2)) if size_match else ""
            prof_sys = prof_match.group(1).strip() if prof_match else ""
            glass = glass_match.group(1).strip() if glass_match else ""
            qty_val = int(qty_match.group(1)) if qty_match else 1

            matched_img_data = page_images[0] if page_images else None

            if base_code:
                if qty_val > 1:
                    for q in range(qty_val):
                        suffix = get_code_suffix(q)
                        split_code = f"{base_code}{suffix}"
                        extracted_items.append({
                            "code": split_code,
                            "name": name,
                            "location": location,
                            "w_val": w_val,
                            "h_val": h_val,
                            "prof_sys": prof_sys,
                            "glass": glass,
                            "mesh": "NO",
                            "lock": "MULTI POINT",
                            "remarks": "",
                            "image_data": matched_img_data,
                        })
                else:
                    extracted_items.append({
                        "code": base_code,
                        "name": name,
                        "location": location,
                        "w_val": w_val,
                        "h_val": h_val,
                        "prof_sys": prof_sys,
                        "glass": glass,
                        "mesh": "NO",
                        "lock": "MULTI POINT",
                        "remarks": "",
                        "image_data": matched_img_data,
                    })

    return extracted_items


def parse_excel_quotation(file_bytes: io.BytesIO) -> List[Dict[str, Any]]:
    wb_source = openpyxl.load_workbook(file_bytes, data_only=True)
    extracted_items = []

    for sheetname in wb_source.sheetnames:
        src_ws = wb_source[sheetname]
        images = getattr(src_ws, "_images", [])

        for r in range(1, src_ws.max_row + 1):
            c1_val = str(src_ws.cell(r, 1).value or "").strip()

            if c1_val == "Code :":
                base_code = str(src_ws.cell(r, 5).value or "").strip()
                size_str = str(
                    src_ws.cell(r, 12).value or src_ws.cell(r, 13).value or ""
                )
                name = str(src_ws.cell(r + 1, 5).value or "").strip()
                prof_sys = str(
                    src_ws.cell(r + 1, 12).value or src_ws.cell(r + 1, 13).value or ""
                ).strip()
                location = str(src_ws.cell(r + 2, 5).value or "").strip()
                glass = str(
                    src_ws.cell(r + 2, 12).value or src_ws.cell(r + 2, 13).value or ""
                ).strip()

                w_val, h_val = "", ""
                if "W =" in size_str and ";" in size_str:
                    try:
                        parts = size_str.split(";")
                        w_val = float(parts[0].split("=")[1].strip())
                        h_val = float(parts[1].split("=")[1].strip())
                    except Exception:
                        pass

                qty_val = 1
                for search_qty_r in range(r, min(r + 20, src_ws.max_row + 1)):
                    lbl_qty = str(src_ws.cell(search_qty_r, 8).value or "").strip()
                    if lbl_qty == "Quantity":
                        try:
                            qty_val = int(
                                src_ws.cell(search_qty_r, 11).value
                                or src_ws.cell(search_qty_r, 12).value
                                or 1
                            )
                        except Exception:
                            qty_val = 1
                        break

                mesh_type, lock_type, remarks_desc = "NO", "", ""

                for offset in range(10, 22):
                    v_mesh = str(src_ws.cell(r + offset, 8).value or "")
                    if "MeshType :" in v_mesh:
                        mesh_type = v_mesh.replace("MeshType :", "").strip() or "NO"

                    v_lock = str(
                        src_ws.cell(r + offset, 13).value
                        or src_ws.cell(r + offset, 14).value
                        or ""
                    )
                    if "Locking :" in v_lock:
                        lock_type = v_lock.replace("Locking :", "").strip()

                for search_r in range(r, min(r + 25, src_ws.max_row + 1)):
                    cell_rem = str(src_ws.cell(search_r, 8).value or "")
                    if "Remarks :" in cell_rem:
                        remarks_desc = cell_rem.replace("Remarks :", "").strip()

                matched_img_data = None
                for img in images:
                    anchor = img.anchor
                    if hasattr(anchor, "_from"):
                        img_row = anchor._from.row + 1
                        if r <= img_row <= r + 22 and img_row > 5:
                            matched_img_data = img._data()
                            break

                if qty_val > 1:
                    for q in range(qty_val):
                        suffix = get_code_suffix(q)
                        split_code = f"{base_code}{suffix}"
                        extracted_items.append({
                            "code": split_code,
                            "name": name,
                            "location": location,
                            "w_val": w_val,
                            "h_val": h_val,
                            "prof_sys": prof_sys,
                            "glass": glass,
                            "mesh": mesh_type,
                            "lock": lock_type,
                            "remarks": remarks_desc,
                            "image_data": matched_img_data,
                        })
                else:
                    extracted_items.append({
                        "code": base_code,
                        "name": name,
                        "location": location,
                        "w_val": w_val,
                        "h_val": h_val,
                        "prof_sys": prof_sys,
                        "glass": glass,
                        "mesh": mesh_type,
                        "lock": lock_type,
                        "remarks": remarks_desc,
                        "image_data": matched_img_data,
                    })

    return extracted_items


def build_measurement_excel(items: List[Dict[str, Any]]) -> Tuple[io.BytesIO, int]:
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "MEASUREMENT"
    ws_out.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(
        start_color="A6A085", end_color="A6A085", fill_type="solid"
    )
    font_header = Font(name="Arial", size=10, bold=True, color="000000")
    font_data = Font(name="Arial", size=10, bold=False, color="000000")

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    align_center = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws_out.row_dimensions[1].height = 28
    ws_out.row_dimensions[2].height = 20

    headers_m = [
        ("A1:A2", "S.NO"),
        ("B1:B2", "FLAT NO"),
        ("C1:C2", "WINDOW CODE"),
        ("D1:D2", "WINDOW TYPE"),
        ("E1:E2", "LOCATION"),
        ("F1:F2", "PROFILE"),
        ("G1:H1", "Dimension in mm"),
        ("I1:I2", "Area in Sq. Ft"),
        ("J1:J2", "Mosquito Mesh"),
        ("K1:K2", "Lock Type"),
        ("L1:M1", "Glass / FLYMESH"),
        ("N1:N2", "Remarks"),
    ]

    for cell_range, text in headers_m:
        if ":" in cell_range:
            ws_out.merge_cells(cell_range)
            ws_out[cell_range.split(":")[0]].value = text
        else:
            ws_out[cell_range].value = text

    ws_out["G2"] = "Width"
    ws_out["H2"] = "Height"
    ws_out["L2"] = "Thickness"
    ws_out["M2"] = "Spece."

    for r in range(1, 3):
        for c in range(1, 15):
            cell = ws_out.cell(row=r, column=c)
            cell.fill = header_fill
            cell.font = font_header
            cell.alignment = align_center
            cell.border = thin_border

    curr_row = 3
    s_no = 1

    for item in items:
        ws_out.row_dimensions[curr_row].height = 24

        ws_out.cell(row=curr_row, column=1, value=s_no)
        ws_out.cell(row=curr_row, column=2, value="")
        ws_out.cell(row=curr_row, column=3, value=item["code"])
        ws_out.cell(row=curr_row, column=4, value=item["name"])
        ws_out.cell(row=curr_row, column=5, value=item["location"])
        ws_out.cell(row=curr_row, column=6, value=item["prof_sys"])
        ws_out.cell(row=curr_row, column=7, value=item["w_val"])
        ws_out.cell(row=curr_row, column=8, value=item["h_val"])

        c_sq = ws_out.cell(
            row=curr_row,
            column=9,
            value=f"=ROUND((G{curr_row}/304.8)*(H{curr_row}/304.8), 2)",
        )
        c_sq.number_format = "0.00"

        glass_thick = extract_glass_thickness(item["glass"])

        ws_out.cell(row=curr_row, column=10, value=item["mesh"])
        ws_out.cell(row=curr_row, column=11, value=item["lock"])
        ws_out.cell(row=curr_row, column=12, value=glass_thick)
        ws_out.cell(row=curr_row, column=13, value=item["glass"])
        ws_out.cell(row=curr_row, column=14, value=item["remarks"])

        for c in range(1, 15):
            cell = ws_out.cell(row=curr_row, column=c)
            cell.font = font_data
            cell.border = thin_border
            if c in [1, 2, 3, 7, 8, 9, 10, 11, 12]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

        # =========================================================
        # AUTOMATIC DRAWING RESIZING & ALIGNMENT (IMAGE 2 STYLE)
        # =========================================================
        if item["image_data"]:
            try:
                img_bytes = io.BytesIO(item["image_data"])
                img_copy = openpyxl.drawing.image.Image(img_bytes)

                # Set clean dimensions to match Image 2 (Width ~ 280px, Height ~ 260px)
                max_width = 280
                max_height = 260

                orig_w, orig_h = img_copy.width, img_copy.height
                if orig_w > 0 and orig_h > 0:
                    ratio = min(max_width / orig_w, max_height / orig_h)
                    img_copy.width = int(orig_w * ratio)
                    img_copy.height = int(orig_h * ratio)
                else:
                    img_copy.width = max_width
                    img_copy.height = max_height

                ws_out.add_image(img_copy, f"C{curr_row + 2}")
            except Exception:
                pass

        # Give 28 rows gap so scaled image fits cleanly below row
        curr_row += 28
        s_no += 1

    col_widths = {
        "A": 6,
        "B": 10,
        "C": 16,
        "D": 22,
        "E": 18,
        "F": 18,
        "G": 10,
        "H": 10,
        "I": 14,
        "J": 15,
        "K": 15,
        "L": 12,
        "M": 25,
        "N": 25,
    }
    for col_letter, width in col_widths.items():
        ws_out.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb_out.save(output)
    output.seek(0)
    return output, len(items)


# ============================================================
# MAIN UI CONTROLLER
# ============================================================

st.markdown(
    "<div class='step-title'>📐 Upload Quotation File (.xlsx / .pdf)</div>",
    unsafe_allow_html=True,
)

uploaded_file = st.file_uploader(
    "Upload File (.xlsx, .xlsm, .pdf)",
    type=["xlsx", "xlsm", "pdf"],
    key="oc_prep_uploader",
)

if st.button("🚀 Generate MEASUREMENT Sheet", type="primary"):
    if not uploaded_file:
        st.error("Please upload an Excel or PDF Quotation file first!")
    else:
        file_bytes = io.BytesIO(uploaded_file.read())
        uploaded_file.seek(0)

        with st.spinner("Processing file, resizing drawings & formatting MEASUREMENT sheet..."):
            if uploaded_file.name.endswith(".pdf"):
                extracted_items = parse_pdf_quotation(file_bytes)
            else:
                extracted_items = parse_excel_quotation(file_bytes)

            if extracted_items:
                excel_out, count = build_measurement_excel(extracted_items)
                st.success(
                    f"Successfully generated MEASUREMENT sheet with perfectly aligned drawings ({count} items split)!"
                )

                out_filename = (
                    f"MEASUREMENT_{uploaded_file.name.rsplit('.', 1)[0]}.xlsx"
                )
                st.download_button(
                    label="📥 Download MEASUREMENT Excel Sheet",
                    data=excel_out,
                    file_name=out_filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=False,
                )
            else:
                st.warning("No valid Window items found in the uploaded file.")
